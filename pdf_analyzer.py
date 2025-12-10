# core/pdf_analyzer.py
"""
PDF 파서 + 계산 규칙을 이용해
성명별 집계를 만들고, 지급조서 템플릿 엑셀에 결과를 채워넣는 모듈.
"""

from __future__ import annotations
from typing import Tuple
import io

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from .pdf_parser import parse_pdf_to_rows
from .rules import compute_allowance_by_person


START_ROW = 5   # 템플릿에서 데이터 시작 행 (연번/직급/성명 있는 줄, C열이 성명)


def summarize_pdf_by_person(df_rows: pd.DataFrame) -> pd.DataFrame:
    """
    행단위 df_rows (parse_pdf_to_rows 결과)를 받아
    PDF 기준 성명별 집계 DataFrame 생성.

    반환:
        index: 성명
        columns:
          - 총지급액_숫자
          - 4시간미만
          - 4시간이상
          - 차량사용횟수
    """
    df = df_rows.copy()

    df["is_under4"] = (df["minutes"] > 0) & (df["minutes"] < 240)
    df["is_over4"] = df["minutes"] >= 240
    df["car_cnt"] = df["car_used"].astype(int)

    grouped = df.groupby("성명").agg(
        총지급액_숫자=("amount_pdf", "sum"),
        _4시간미만=("is_under4", "sum"),
        _4시간이상=("is_over4", "sum"),
        차량사용횟수=("car_cnt", "sum"),
    )

    grouped = grouped.rename(columns={"_4시간미만": "4시간미만", "_4시간이상": "4시간이상"})
    for c in ["총지급액_숫자", "4시간미만", "4시간이상", "차량사용횟수"]:
        grouped[c] = grouped[c].fillna(0).astype(int)

    return grouped


def fill_template_with_summary(template_bytes: bytes, summary: pd.DataFrame) -> bytes:
    """
    지급조서 템플릿(엑셀 바이너리)에 summary 정보를 채워넣고,
    결과 엑셀을 bytes로 반환.

    템플릿 가정:
      - C열: 성명
      - D열: 4시간 미만 횟수
      - E열: 4시간 이상 횟수
      - F열: 차량 사용 횟수
      - H열: PDF 기준 지급액
      - L열: 규칙에 따른 계산금액 (PDF 금액과 다를 때만 표기)
      - H열과 L열 값이 다르면 H열은 빨간색/굵게 표시
    """
    summary = summary.copy()
    summary = summary.fillna(0)
    # 타입 보정
    for col in ["총지급액_숫자", "4시간미만", "4시간이상", "차이", "차량사용횟수", "계산_총지급액"]:
        if col in summary.columns:
            summary[col] = summary[col].astype(int)

    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb.active

    header_row = START_ROW - 1
    # L열 헤더 없으면 추가 (선택)
    if not ws[f"L{header_row}"].value:
        ws[f"L{header_row}"] = "계산금액(규칙)"

    row = START_ROW
    last_filled_row = START_ROW - 1

    while True:
        cell_name = ws[f"C{row}"]
        name_val = cell_name.value

        if not name_val or str(name_val).strip() in ["합계"]:
            break

        name_str = str(name_val).strip()

        # 기본값
        under4 = over4 = car_cnt = 0
        amount_pdf = amount_calc = 0

        if name_str in summary.index:
            r = summary.loc[name_str]
            under4 = int(r.get("4시간미만", 0) or 0)
            over4 = int(r.get("4시간이상", 0) or 0)
            car_cnt = int(r.get("차량사용횟수", 0) or 0)
            amount_pdf = int(r.get("총지급액_숫자", 0) or 0)
            amount_calc = int(r.get("계산_총지급액", 0) or 0)

        # D/E/F: 4시간 미만/이상/차량횟수
        ws[f"D{row}"] = "" if under4 == 0 else under4
        ws[f"E{row}"] = "" if over4 == 0 else over4
        ws[f"F{row}"] = "" if car_cnt == 0 else car_cnt

        # H열: PDF 기준 지급액
        cell_h = ws[f"H{row}"]
        cell_h.value = "" if amount_pdf == 0 else amount_pdf

        # L열: 계산금액 (다를 때만 표시)
        cell_l = ws[f"L{row}"]
        if amount_pdf != amount_calc:
            cell_l.value = amount_calc
            # H열을 빨간색/굵게
            cell_h.font = Font(color="FF0000", bold=True)
        else:
            cell_l.value = ""

        last_filled_row = row
        row += 1

    # 합계 행(H열) 자동 계산 (원하면 사용, 아니면 제거 가능)
    if last_filled_row >= START_ROW:
        total_row = last_filled_row + 1
        if not ws[f"G{total_row}"].value:
            ws[f"G{total_row}"] = "합계"
        ws[f"H{total_row}"] = f"=SUM(H{START_ROW}:H{last_filled_row})"

    # bytes로 저장
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def analyze_pdf_and_template(
    pdf_bytes: bytes,
    template_bytes: bytes,
) -> Tuple[pd.DataFrame, bytes]:
    """
    PDF + 템플릿을 받아서:
      - 성명별 요약 DataFrame(summary)
      - 템플릿에 결과 채운 엑셀 bytes
    를 반환.
    """
    # 1) PDF 파싱
    df_rows = parse_pdf_to_rows(pdf_bytes)

    # 2) PDF 기준 집계
    pdf_summary = summarize_pdf_by_person(df_rows)

    # 3) 규칙에 따른 실제 지급액 계산
    calc_totals = compute_allowance_by_person(df_rows)

    # 4) summary 결합
    summary = pdf_summary.copy()
    summary["계산_총지급액"] = calc_totals
    summary["계산_총지급액"] = summary["계산_총지급액"].fillna(0).astype(int)
    summary["차이"] = summary["계산_총지급액"] - summary["총지급액_숫자"]

    # 5) 템플릿 채우기
    result_bytes = fill_template_with_summary(template_bytes, summary)

    return summary, result_bytes
